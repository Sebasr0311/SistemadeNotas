"""
Tests de la configuración de cálculo (ColumnConfig): columnas seleccionables
y promedios con pesos.

Cubre:
- Creación de la config por defecto desde una planilla.
- Fórmulas del modo simple con columnas seleccionadas.
- Fórmulas del modo pesos con pesos que suman 100 y que no suman 100.
- Generación del Excel end-to-end con ambas configuraciones.
- Que el modo legacy (sin ColumnConfig) sigue generando la fórmula histórica.

Uso (desde la raíz del proyecto):
    python -m tests.test_generar_excel_columnas_config
"""

import os
import sys
import tempfile
import unittest

# La raíz del proyecto se agrega al path para poder importar el paquete excel/.
_PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)

import openpyxl  # noqa: E402

import excel.generar_excel_notas as generador  # noqa: E402
from excel.generar_excel_notas import ColumnConfig  # noqa: E402


def _planilla(area_por_alumno, n_area_trabajo=None, periodo=3):
    """Construye una planilla de prueba (misma forma que en los otros tests)."""
    encabezado = {
        "institucion": "INSTITUCION DEMO",
        "sede": "SEDE",
        "año_lectivo": "2026",
        "jornada": "MAÑANA",
        "grupo": "0302",
        "asignatura": "MATEMATICAS",
        "docente": "DOCENTE DEMO",
        "periodo": periodo,
    }
    if n_area_trabajo is not None:
        encabezado["n_area_trabajo"] = n_area_trabajo

    estudiantes = []
    for i, notas in enumerate(area_por_alumno, start=1):
        retirado = notas is None
        estudiantes.append({
            "no": i,
            "nombre": f"ALUMNO {i}",
            "ev_anteriores": [45, 45] if periodo > 1 else [],
            "area_trabajo": notas if not retirado else None,
            "retirado": retirado,
            "revisar": [False] * (len(notas) if not retirado else 0),
        })
    return {"encabezado": encabezado, "estudiantes": estudiantes}


def _hoja_cargada(planilla, column_config):
    """Genera el Excel de la planilla con la config y devuelve la hoja activa."""
    directorio = tempfile.mkdtemp(prefix="notas_config_")
    ruta = os.path.join(directorio, "salida.xlsx")
    generador.generar_excel_planilla(planilla, ruta, column_config=column_config)
    wb = openpyxl.load_workbook(ruta)
    return wb[wb.sheetnames[0]]


def _fila_header(ws):
    """Ubica la fila donde empieza la tabla (la que tiene 'No.' y 'Nombre')."""
    for r in range(1, 20):
        if ws.cell(row=r, column=1).value == "No." and \
           ws.cell(row=r, column=2).value == "Nombre del Alumno":
            return r
    raise AssertionError("No se encontró la fila de encabezado de la tabla")


class TestColumnConfigBase(unittest.TestCase):
    """ColumnConfig: creación, selección y validación básica."""

    def test_crear_desde_planilla_simple_todas_incluidas(self):
        cfg = ColumnConfig.crear_desde_planilla(4)
        self.assertEqual(cfg.modo, "simple")
        self.assertEqual(len(cfg.columnas), 4)
        self.assertTrue(all(c["incluida"] for c in cfg.columnas))
        self.assertEqual(cfg.columnas_seleccionadas, [0, 1, 2, 3])
        self.assertEqual(cfg.n_columnas_seleccionadas, 4)
        # Peso por defecto reparte 100 en partes iguales
        self.assertEqual(cfg.columnas[0]["peso"], 25.0)

    def test_es_pesado_solo_con_mas_de_una_columna(self):
        cfg = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "A1", "incluida": True, "peso": 100},
        ])
        # Una sola columna: no es un promedio pesado real (usa el valor directo)
        self.assertFalse(cfg.es_pesado())

        cfg2 = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "A1", "incluida": True, "peso": 60},
            {"nombre": "A2", "incluida": True, "peso": 40},
        ])
        self.assertTrue(cfg2.es_pesado())

    def test_pesos_suman_cien_con_tolerancia(self):
        cfg = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "A1", "incluida": True, "peso": 33.3},
            {"nombre": "A2", "incluida": True, "peso": 33.3},
            {"nombre": "A3", "incluida": True, "peso": 33.4},
        ])
        self.assertTrue(cfg.pesos_suman_cien())

    def test_pesos_no_suman_cien(self):
        cfg = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "A1", "incluida": True, "peso": 30},
            {"nombre": "A2", "incluida": True, "peso": 30},
        ])
        self.assertFalse(cfg.pesos_suman_cien())

    def test_roundtrip_dict(self):
        cfg = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "A1", "incluida": True, "peso": 60},
            {"nombre": "A2", "incluida": False, "peso": 0},
        ])
        cfg2 = ColumnConfig.from_dict(cfg.to_dict())
        self.assertEqual(cfg2.modo, "pesos")
        self.assertEqual(cfg2.columnas, cfg.columnas)

    def test_modo_invalido_cae_a_simple(self):
        cfg = ColumnConfig(modo="raro")
        self.assertEqual(cfg.modo, "simple")

    def test_crear_desde_planilla_con_ev(self):
        # Periodo 3 => 2 evs + 2 áreas: primero "Def. Periodo 1..2", luego
        # "Área Trabajo 1..2", cada una con su tipo y pos.
        cfg = ColumnConfig.crear_desde_planilla(2, n_ev=2)
        self.assertEqual(len(cfg.columnas), 4)
        self.assertEqual(cfg.columnas[0]["nombre"], "Def. Periodo 1")
        self.assertEqual(cfg.columnas[0]["tipo"], "ev")
        self.assertEqual(cfg.columnas[0]["pos"], 0)
        self.assertEqual(cfg.columnas[1]["nombre"], "Def. Periodo 2")
        self.assertEqual(cfg.columnas[1]["pos"], 1)
        self.assertEqual(cfg.columnas[2]["nombre"], "Área Trabajo 1")
        self.assertEqual(cfg.columnas[2]["tipo"], "area")
        self.assertEqual(cfg.columnas[2]["pos"], 0)
        self.assertEqual(cfg.columnas[3]["nombre"], "Área Trabajo 2")
        self.assertEqual(cfg.columnas[3]["pos"], 1)
        # Peso por defecto reparte 100 entre el total (4).
        self.assertEqual(cfg.columnas[0]["peso"], 25.0)
        self.assertTrue(all(c["incluida"] for c in cfg.columnas))
        self.assertEqual(cfg.columnas_seleccionadas, [0, 1, 2, 3])


class TestFormulasUnicas(unittest.TestCase):
    """_formula_definitiva: las fórmulas según la configuración."""

    def _f(self, column_config):
        # Periodo 3 => ev previas en C (3) y D (4), área arranca en E (col_at1=5).
        # 5 columnas de área => col_def = 5+5 = 10 (col_ev_start=3, col_at1=5).
        return generador._formula_definitiva(3, 5, 11, 5, column_config)

    def test_legacy_sin_config_sum_todas(self):
        # Sin config: SUM de las áreas (E..I) / n_areas.
        self.assertEqual(
            self._f(None),
            '=IFERROR(SUM(E11:I11)/5,"")',
        )

    def test_simple_con_todas_average_rango_completo(self):
        cfg = ColumnConfig.crear_desde_planilla(5)
        self.assertEqual(
            self._f(cfg),
            '=IFERROR(AVERAGE(E11,F11,G11,H11,I11),"")',
        )

    def test_simple_solo_columnas_seleccionadas(self):
        # Solo columnas 1 y 3 (0-based: índices 0 y 2) -> E y G
        cfg = ColumnConfig(
            columnas=[{"nombre": "A1", "incluida": True, "peso": 50},
                      {"nombre": "A2", "incluida": False, "peso": 0},
                      {"nombre": "A3", "incluida": True, "peso": 50},
                      {"nombre": "A4", "incluida": False, "peso": 0},
                      {"nombre": "A5", "incluida": False, "peso": 0}],
        )
        self.assertEqual(
            self._f(cfg),
            '=IFERROR(AVERAGE(E11,G11),"")',
        )

    def test_simple_una_sola_columna_seleccionada_referencia_directa(self):
        cfg = ColumnConfig(columnas=[
            {"nombre": "A1", "incluida": True, "peso": 100},
            {"nombre": "A2", "incluida": False, "peso": 0},
            {"nombre": "A3", "incluida": False, "peso": 0},
            {"nombre": "A4", "incluida": False, "peso": 0},
            {"nombre": "A5", "incluida": False, "peso": 0},
        ])
        self.assertEqual(self._f(cfg), '=IFERROR(E11,"")')

    def test_ninguna_columna_seleccionada(self):
        cfg = ColumnConfig(columnas=[
            {"nombre": "A1", "incluida": False, "peso": 0},
            {"nombre": "A2", "incluida": False, "peso": 0},
        ])
        self.assertEqual(self._f(cfg), "")

    def test_pesos_suman_cien(self):
        cfg = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "A1", "incluida": True, "peso": 60},
            {"nombre": "A2", "incluida": True, "peso": 40},
            {"nombre": "A3", "incluida": False, "peso": 0},
            {"nombre": "A4", "incluida": False, "peso": 0},
            {"nombre": "A5", "incluida": False, "peso": 0},
        ])
        # 60*E + 40*F, todo /100
        self.assertEqual(
            self._f(cfg),
            '=IFERROR((E11*60+F11*40)/100,"")',
        )

    def test_pesos_no_suman_cien_se_normaliza(self):
        cfg = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "A1", "incluida": True, "peso": 30},
            {"nombre": "A2", "incluida": True, "peso": 30},
            {"nombre": "A3", "incluida": False, "peso": 0},
            {"nombre": "A4", "incluida": False, "peso": 0},
            {"nombre": "A5", "incluida": False, "peso": 0},
        ])
        # 30*E + 30*F, todo /(30+30) — así la nota queda en la misma escala.
        self.assertEqual(
            self._f(cfg),
            '=IFERROR((E11*30+F11*30)/(30+30),"")',
        )

    def test_pesos_una_columna_ref_directa(self):
        cfg = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "A1", "incluida": True, "peso": 100},
            {"nombre": "A2", "incluida": False, "peso": 0},
        ])
        self.assertEqual(self._f(cfg), '=IFERROR(E11,"")')

    # ── Columnas mixtas (Def. Periodo + Área Trabajo) ──
    # Periodo 3 => col_ev_start=3 (Def. Periodo 1 en C, Def. Periodo 2 en D),
    # col_at1=5 (Área Trabajo 1 en E, Área Trabajo 2 en F). 5 áreas => col_def 10.
    def _cfg_mixta(self, seleccion):
        """2 evs + 2 áreas: [Def.P1, Def.P2, Área1, Área2]."""
        columnas = [
            {"nombre": "Def. Periodo 1", "tipo": "ev", "pos": 0, "incluida": "D1" in seleccion, "peso": 25},
            {"nombre": "Def. Periodo 2", "tipo": "ev", "pos": 1, "incluida": "D2" in seleccion, "peso": 25},
            {"nombre": "Área Trabajo 1", "tipo": "area", "pos": 0, "incluida": "A1" in seleccion, "peso": 25},
            {"nombre": "Área Trabajo 2", "tipo": "area", "pos": 1, "incluida": "A2" in seleccion, "peso": 25},
        ]
        return ColumnConfig(columnas=[dict(c, incluida=bool(c["incluida"])) for c in columnas])

    def _f_mixta_simple(self, seleccion):
        # La misma base física de self._f: col_ev_start=3, col_at1=5, r=11.
        return generador._formula_definitiva(3, 5, 11, 5, self._cfg_mixta(seleccion))

    def test_simple_mixto_def_periodo1_y_area2(self):
        # Def. Periodo 1 -> C(3); Área Trabajo 2 -> F(6) (col_at1=5 + pos 1).
        self.assertEqual(
            self._f_mixta_simple({"D1", "A2"}),
            '=IFERROR(AVERAGE(C11,F11),"")',
        )

    def test_pesos_mixto_def_periodo_y_area(self):
        cfg = self._cfg_mixta({"D1", "D2", "A1", "A2"})
        cfg.modo = "pesos"
        for c in cfg.columnas:
            c["peso"] = 25
        self.assertEqual(
            generador._formula_definitiva(3, 5, 11, 5, cfg),
            '=IFERROR((C11*25+D11*25+E11*25+F11*25)/100,"")',
        )


class TestExcelEndToEndConConfig(unittest.TestCase):
    """El Excel final refleja la configuración elegida."""

    def test_simple_con_columnas_seleccionadas(self):
        planilla = _planilla(
            [[40, 50, 60, 70], [41, 51, 61, 71]],
            n_area_trabajo=4,
        )
        cfg = ColumnConfig(columnas=[
            {"nombre": "Área Trabajo 1", "incluida": True, "peso": 25},
            {"nombre": "Área Trabajo 2", "incluida": False, "peso": 0},
            {"nombre": "Área Trabajo 3", "incluida": True, "peso": 25},
            {"nombre": "Área Trabajo 4", "incluida": False, "peso": 0},
        ])
        ws = _hoja_cargada(planilla, cfg)
        # Periodo 3 -> ev previas en C y D; área arranca en E(5)..H(8);
        # definitiva en I(9). Alumno 1 = fila header + 1.
        hr = _fila_header(ws)
        fila1 = hr + 1
        # Solo columnas 1 y 3 (0-based: índices 0 y 2) -> E y G
        self.assertEqual(ws.cell(row=fila1, column=9).value,
                         '=IFERROR(AVERAGE(E{f},G{f}),"")'.format(f=fila1))
        # La columna excluida conserva su nota, pero no entra a la fórmula.
        self.assertEqual(ws.cell(row=fila1, column=6).value, 50)

    def test_pesos_end_to_end(self):
        planilla = _planilla(
            [[40, 50], [41, 51]],
            n_area_trabajo=2,
        )
        cfg = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "Área Trabajo 1", "incluida": True, "peso": 60},
            {"nombre": "Área Trabajo 2", "incluida": True, "peso": 40},
        ])
        ws = _hoja_cargada(planilla, cfg)
        # Periodo 3 -> ev previas en C y D; área E(5),F(6); definitiva en G(7).
        hr = _fila_header(ws)
        fila1 = hr + 1
        self.assertEqual(ws.cell(row=fila1, column=7).value,
                         '=IFERROR((E{f}*60+F{f}*40)/100,"")'.format(f=fila1))

        # La info del Excel indica el modo de cálculo y los pesos.
        valores_info = [ws.cell(row=r, column=1).value for r in range(1, hr)]
        self.assertIn("Modo de cálculo", valores_info)
        self.assertIn("Pesos", valores_info)

    def test_legacy_sigue_generando_formula_historica(self):
        planilla = _planilla(
            [[40, 50], [41, 51]],
            n_area_trabajo=2,
        )
        ws = _hoja_cargada(planilla, None)
        # Sin config: la fórmula histórica SUM/2 (periodo 3 -> área en E,F,
        # definitiva en G). Sin info extra, header en 10, alumno 1 en 11.
        self.assertEqual(ws.cell(row=11, column=7).value,
                         '=IFERROR(SUM(E11:F11)/2,"")')

    def test_mixto_end_to_end_simple(self):
        # Planilla periodo 3 (2 evs) con 2 áreas. Selección mixta: "Def. Periodo 1"
        # (col C = 3) y "Área Trabajo 2" (col F = 6). En modo simple => AVERAGE.
        planilla = _planilla(
            [[40, 50], [41, 51]],
            n_area_trabajo=2,
        )
        cfg = ColumnConfig(columnas=[
            {"nombre": "Def. Periodo 1", "tipo": "ev", "pos": 0, "incluida": True, "peso": 0},
            {"nombre": "Def. Periodo 2", "tipo": "ev", "pos": 1, "incluida": False, "peso": 0},
            {"nombre": "Área Trabajo 1", "tipo": "area", "pos": 0, "incluida": False, "peso": 0},
            {"nombre": "Área Trabajo 2", "tipo": "area", "pos": 1, "incluida": True, "peso": 0},
        ])
        # ev en C(3),D(4); área en E(5),F(6); definitiva en G(7).
        ws = _hoja_cargada(planilla, cfg)
        hr = _fila_header(ws)
        fila1 = hr + 1
        self.assertEqual(ws.cell(row=fila1, column=7).value,
                         f'=IFERROR(AVERAGE(C{fila1},F{fila1}),"")')
        # Valores físicos: ev1 en C, área2 en F.
        self.assertEqual(ws.cell(row=fila1, column=3).value, 45)  # ev1
        self.assertEqual(ws.cell(row=fila1, column=6).value, 50)  # área2

    def test_mixto_end_to_end_pesos(self):
        # Mismo escenario, modo pesos con Def. Periodo 1 (60) y Área Trabajo 2 (40).
        planilla = _planilla(
            [[40, 50], [41, 51]],
            n_area_trabajo=2,
        )
        cfg = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "Def. Periodo 1", "tipo": "ev", "pos": 0, "incluida": True, "peso": 60},
            {"nombre": "Def. Periodo 2", "tipo": "ev", "pos": 1, "incluida": False, "peso": 0},
            {"nombre": "Área Trabajo 1", "tipo": "area", "pos": 0, "incluida": False, "peso": 0},
            {"nombre": "Área Trabajo 2", "tipo": "area", "pos": 1, "incluida": True, "peso": 40},
        ])
        ws = _hoja_cargada(planilla, cfg)
        hr = _fila_header(ws)
        fila1 = hr + 1
        self.assertEqual(ws.cell(row=fila1, column=7).value,
                         f'=IFERROR((C{fila1}*60+F{fila1}*40)/100,"")')


class TestExcelPorFormaConColumnConfigs(unittest.TestCase):
    """La config se resuelve POR FORMA cuando se pasa `column_configs`:
    cada hoja usa la config de su propia cantidad de columnas (spec v3)."""

    def _planillas(self):
        # Forma n=2 (grupo 0201) y forma n=4 (grupo 0401), mismo periodo 3.
        p2 = _planilla([[40, 50], [41, 51]], n_area_trabajo=2)
        p2["encabezado"]["grupo"] = "0201"
        p4 = _planilla([[40, 50, 60, 70], [41, 51, 61, 71]], n_area_trabajo=4)
        p4["encabezado"]["grupo"] = "0401"
        return [p2, p4]

    def _configs(self):
        cfg2 = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "Área Trabajo 1", "incluida": True, "peso": 60},
            {"nombre": "Área Trabajo 2", "incluida": True, "peso": 40},
        ])
        cfg4 = ColumnConfig(modo="simple", columnas=[
            {"nombre": "Área Trabajo 1", "incluida": True, "peso": 25},
            {"nombre": "Área Trabajo 2", "incluida": True, "peso": 25},
            {"nombre": "Área Trabajo 3", "incluida": True, "peso": 25},
            {"nombre": "Área Trabajo 4", "incluida": True, "peso": 25},
        ])
        return {("n_areas", 2): cfg2, ("n_areas", 4): cfg4}

    def test_cada_hoja_usa_la_config_de_su_forma(self):
        directorio = tempfile.mkdtemp(prefix="notas_por_forma_")
        ruta = os.path.join(directorio, "salida.xlsx")
        generador.generar_excel_asignatura(
            self._planillas(), ruta, column_configs=self._configs()
        )

        wb = openpyxl.load_workbook(ruta)
        ws2 = wb["Curso 0201 - MATEMATICAS"]
        ws4 = wb["Curso 0401 - MATEMATICAS"]

        # Forma n=2 (pesos 60/40): área E,F; definitiva en G(7).
        hr2 = _fila_header(ws2)
        fila1 = hr2 + 1
        self.assertEqual(ws2.cell(row=fila1, column=7).value,
                         '=IFERROR((E{f}*60+F{f}*40)/100,"")'.format(f=fila1))
        # La info de la hoja refleja el modo y los pesos.
        valores2 = [ws2.cell(row=r, column=1).value for r in range(1, hr2)]
        self.assertIn("Modo de cálculo", valores2)
        self.assertIn("Pesos", valores2)
        self.assertIn("Pesos", [ws2.cell(row=r, column=2).value for r in range(1, hr2)])

        # Forma n=4 (simple con todas): área E..H; definitiva en I(9).
        hr4 = _fila_header(ws4)
        fila1 = hr4 + 1
        self.assertEqual(ws4.cell(row=fila1, column=9).value,
                         '=IFERROR(AVERAGE(E{f},F{f},G{f},H{f}),"")'.format(f=fila1))
        valores4 = [ws4.cell(row=r, column=1).value for r in range(1, hr4)]
        self.assertIn("Modo de cálculo", valores4)
        self.assertNotIn("Pesos", valores4)
        self.assertIn("Promedio simple", [ws4.cell(row=r, column=2).value for r in range(1, hr4)])

    def test_sin_clave_para_una_forma_cae_a_legacy(self):
        # column_configs solo tiene la forma n=2: la hoja n=4 queda legacy
        # (fórmula SUM/n, sin filas de info de cálculo).
        directorio = tempfile.mkdtemp(prefix="notas_por_forma_")
        ruta = os.path.join(directorio, "salida.xlsx")
        configs = dict(self._configs())
        del configs[("n_areas", 4)]
        generador.generar_excel_asignatura(
            self._planillas(), ruta, column_configs=configs
        )

        wb = openpyxl.load_workbook(ruta)
        ws4 = wb["Curso 0401 - MATEMATICAS"]
        hr4 = _fila_header(ws4)
        fila1 = hr4 + 1
        # Legacy: col_def = 5 + 4 = 9, fórmula SUM(E..H)/4.
        self.assertEqual(ws4.cell(row=fila1, column=9).value,
                         f'=IFERROR(SUM(E{fila1}:H{fila1})/4,"")')
        valores4 = [ws4.cell(row=r, column=1).value for r in range(1, hr4)]
        self.assertNotIn("Modo de cálculo", valores4)


class TestResolverColumnConfigPorIndice(unittest.TestCase):
    """_resolver_column_config: la config se resuelve POR PLANILLA cuando el
    dict `column_configs` está keyeado por índice de planilla (`_idx`).

    Este es el fix del bug: dos planillas con la MISMA cantidad de columnas
    (misma forma) no deben compartir config; cada una usa la suya.
    """

    def test_resolver_por_indice_mismo_forma(self):
        # Dos planillas con 4 columnas (MISMA forma): el resolver por índice
        # las separa y cada una usa su propia config.
        p0 = _planilla([[40, 50, 60, 70]], n_area_trabajo=4)
        p0["_idx"] = 0
        p1 = _planilla([[40, 50, 60, 70]], n_area_trabajo=4)
        p1["_idx"] = 1

        cfg0 = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "Área Trabajo 1", "incluida": True, "peso": 60},
            {"nombre": "Área Trabajo 2", "incluida": True, "peso": 40},
            {"nombre": "Área Trabajo 3", "incluida": False, "peso": 0},
            {"nombre": "Área Trabajo 4", "incluida": False, "peso": 0},
        ])
        cfg1 = ColumnConfig(columnas=[
            {"nombre": "Área Trabajo 1", "incluida": True, "peso": 25},
            {"nombre": "Área Trabajo 2", "incluida": False, "peso": 0},
            {"nombre": "Área Trabajo 3", "incluida": True, "peso": 25},
            {"nombre": "Área Trabajo 4", "incluida": False, "peso": 0},
        ])

        configs = {0: cfg0, 1: cfg1}
        self.assertIs(generador._resolver_column_config(p0, None, configs), cfg0)
        self.assertIs(generador._resolver_column_config(p1, None, configs), cfg1)

    def test_resolver_por_indice_ignora_forma_legacy(self):
        # Planilla con `_idx` presente: aunque el dict también tenga la clave
        # de forma, gana el índice (no la forma).
        p0 = _planilla([[40, 50]], n_area_trabajo=2)
        p0["_idx"] = 0
        cfg_idx = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "Área Trabajo 1", "incluida": True, "peso": 60},
            {"nombre": "Área Trabajo 2", "incluida": True, "peso": 40},
        ])
        cfg_forma = ColumnConfig(columnas=[
            {"nombre": "Área Trabajo 1", "incluida": True, "peso": 50},
            {"nombre": "Área Trabajo 2", "incluida": True, "peso": 50},
        ])
        configs = {0: cfg_idx, ("n_areas", 2): cfg_forma}
        self.assertIs(generador._resolver_column_config(p0, None, configs), cfg_idx)

    def test_resolver_legacy_por_forma_sin_idx(self):
        # Planillas SIN `_idx` (retrocompat): se resuelve por forma.
        p2 = _planilla([[40, 50]], n_area_trabajo=2)
        p4 = _planilla([[40, 50, 60, 70]], n_area_trabajo=4)
        cfg2 = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "Área Trabajo 1", "incluida": True, "peso": 60},
            {"nombre": "Área Trabajo 2", "incluida": True, "peso": 40},
        ])
        cfg4 = ColumnConfig(columnas=[
            {"nombre": "Área Trabajo 1", "incluida": True, "peso": 25},
            {"nombre": "Área Trabajo 2", "incluida": True, "peso": 25},
            {"nombre": "Área Trabajo 3", "incluida": True, "peso": 25},
            {"nombre": "Área Trabajo 4", "incluida": True, "peso": 25},
        ])
        configs = {("n_areas", 2): cfg2, ("n_areas", 4): cfg4}
        self.assertIs(generador._resolver_column_config(p2, None, configs), cfg2)
        self.assertIs(generador._resolver_column_config(p4, None, configs), cfg4)

    def test_resolver_sin_configs_cae_a_column_config(self):
        p = _planilla([[40, 50]], n_area_trabajo=2)
        p["_idx"] = 3
        cfg = ColumnConfig(columnas=[
            {"nombre": "A1", "incluida": True, "peso": 100},
            {"nombre": "A2", "incluida": False, "peso": 0},
        ])
        self.assertIs(generador._resolver_column_config(p, cfg, None), cfg)
        self.assertIsNone(generador._resolver_column_config(p, None, None))


class TestExcelPorIndiceConColumnConfigs(unittest.TestCase):
    """generar_excel_asignatura con `column_configs={indice: ColumnConfig}`:
    cada hoja usa la config de SU planilla, aunque comparta forma con otra."""

    def test_mismo_forma_cada_hoja_usa_la_config_de_su_planilla(self):
        # Dos planillas con 4 columnas (MISMA forma) pero configs distintas:
        # la hoja de cada una debe usar la config de su planilla.
        p0 = _planilla([[40, 50, 60, 70], [41, 51, 61, 71]], n_area_trabajo=4)
        p0["encabezado"]["grupo"] = "0201"
        p0["_idx"] = 0
        p1 = _planilla([[40, 50, 60, 70], [41, 51, 61, 71]], n_area_trabajo=4)
        p1["encabezado"]["grupo"] = "0401"
        p1["_idx"] = 1

        cfg0 = ColumnConfig(modo="pesos", columnas=[
            {"nombre": "Área Trabajo 1", "incluida": True, "peso": 60},
            {"nombre": "Área Trabajo 2", "incluida": True, "peso": 40},
            {"nombre": "Área Trabajo 3", "incluida": False, "peso": 0},
            {"nombre": "Área Trabajo 4", "incluida": False, "peso": 0},
        ])
        cfg1 = ColumnConfig(columnas=[
            {"nombre": "Área Trabajo 1", "incluida": True, "peso": 25},
            {"nombre": "Área Trabajo 2", "incluida": False, "peso": 0},
            {"nombre": "Área Trabajo 3", "incluida": True, "peso": 25},
            {"nombre": "Área Trabajo 4", "incluida": False, "peso": 0},
        ])

        directorio = tempfile.mkdtemp(prefix="notas_por_indice_")
        ruta = os.path.join(directorio, "salida.xlsx")
        generador.generar_excel_asignatura(
            [p0, p1], ruta, column_configs={0: cfg0, 1: cfg1}
        )

        wb = openpyxl.load_workbook(ruta)
        ws0 = wb["Curso 0201 - MATEMATICAS"]
        ws1 = wb["Curso 0401 - MATEMATICAS"]

        # Misma forma (4 áreas E..H, definitiva en I(9)) pero fórmulas
        # distintas: la de cada planilla.
        hr0 = _fila_header(ws0)
        fila0 = hr0 + 1
        self.assertEqual(ws0.cell(row=fila0, column=9).value,
                         '=IFERROR((E{f}*60+F{f}*40)/100,"")'.format(f=fila0))

        hr1 = _fila_header(ws1)
        fila1 = hr1 + 1
        self.assertEqual(ws1.cell(row=fila1, column=9).value,
                         '=IFERROR(AVERAGE(E{f},G{f}),"")'.format(f=fila1))

        # La info de cada hoja refleja su propia config.
        valores0 = [ws0.cell(row=r, column=1).value for r in range(1, hr0)]
        self.assertIn("Modo de cálculo", valores0)
        self.assertIn("Pesos", valores0)
        valores1 = [ws1.cell(row=r, column=1).value for r in range(1, hr1)]
        self.assertIn("Modo de cálculo", valores1)
        self.assertNotIn("Pesos", valores1)


def _planilla_con_otras(area_por_alumno, otras_por_alumno, n_area_trabajo=None,
                        otras_columnas=None, periodo=3):
    """Planilla con columnas "otras" (trabajo práctico, parcial, etc.).

    Misma base que _planilla; `otras_columnas=None` omite la clave del
    encabezado (como las planillas pre-cambio), [] la declara vacía.
    """
    encabezado = {
        "institucion": "INSTITUCION DEMO",
        "sede": "SEDE",
        "año_lectivo": "2026",
        "jornada": "MAÑANA",
        "grupo": "0302",
        "asignatura": "MATEMATICAS",
        "docente": "DOCENTE DEMO",
        "periodo": periodo,
    }
    if n_area_trabajo is not None:
        encabezado["n_area_trabajo"] = n_area_trabajo
    if otras_columnas is not None:
        encabezado["otras_columnas"] = list(otras_columnas)

    estudiantes = []
    for i, notas in enumerate(area_por_alumno, start=1):
        retirado = notas is None
        otras = otras_por_alumno[i - 1] if otras_por_alumno else []
        estudiantes.append({
            "no": i,
            "nombre": f"ALUMNO {i}",
            "ev_anteriores": [45, 45] if periodo > 1 else [],
            "area_trabajo": notas if not retirado else None,
            "otras_notas": [] if retirado else list(otras),
            "retirado": retirado,
            "revisar": [False] * (len(notas) if not retirado else 0),
            "revisar_otras": [False] * (len(otras) if not retirado else 0),
        })
    return {"encabezado": encabezado, "estudiantes": estudiantes}


class TestColumnConfigOtras(unittest.TestCase):
    """ColumnConfig.crear_desde_planilla con columnas "otras"."""

    def test_crear_desde_planilla_con_otras(self):
        # Periodo 3 (2 evs) + 2 áreas + 2 otras: orden ev -> área -> otras.
        cfg = ColumnConfig.crear_desde_planilla(
            2, n_ev=2, nombres_otras=["Parcial", "Trabajo Práctico"]
        )
        self.assertEqual(len(cfg.columnas), 6)
        self.assertEqual(cfg.columnas[0]["nombre"], "Def. Periodo 1")
        self.assertEqual(cfg.columnas[0]["tipo"], "ev")
        self.assertEqual(cfg.columnas[2]["nombre"], "Área Trabajo 1")
        self.assertEqual(cfg.columnas[2]["tipo"], "area")
        self.assertEqual(cfg.columnas[4]["nombre"], "Parcial")
        self.assertEqual(cfg.columnas[4]["tipo"], "otra")
        self.assertEqual(cfg.columnas[4]["pos"], 0)
        self.assertEqual(cfg.columnas[5]["nombre"], "Trabajo Práctico")
        self.assertEqual(cfg.columnas[5]["tipo"], "otra")
        self.assertEqual(cfg.columnas[5]["pos"], 1)
        # Peso por defecto reparte 100 entre el total (6).
        self.assertEqual(cfg.columnas[0]["peso"], round(100.0 / 6, 1))
        self.assertTrue(all(c["incluida"] for c in cfg.columnas))
        self.assertEqual(cfg.columnas_seleccionadas, [0, 1, 2, 3, 4, 5])

    def test_sin_otras_conserva_firma_anterior(self):
        # Llamada con 2 args (retrocompat): misma config que antes.
        cfg = ColumnConfig.crear_desde_planilla(2, n_ev=2)
        self.assertEqual(len(cfg.columnas), 4)
        self.assertEqual([c["tipo"] for c in cfg.columnas],
                         ["ev", "ev", "area", "area"])

    def test_nombre_vacio_usa_generico(self):
        cfg = ColumnConfig.crear_desde_planilla(2, n_ev=0,
                                                nombres_otras=["  ", "Parcial"])
        self.assertEqual(cfg.columnas[2]["nombre"], "Otras notas 1")
        self.assertEqual(cfg.columnas[3]["nombre"], "Parcial")


class TestFormulasConOtras(unittest.TestCase):
    """_formula_definitiva con columnas "otras" y col_otras_start."""

    def _cfg(self, seleccion, modo="simple"):
        """[Def.P1, Def.P2, A1, A2, O1, O2] con pesos 25."""
        columnas = []
        for k in range(2):
            columnas.append({"nombre": f"Def. Periodo {k+1}", "tipo": "ev",
                             "pos": k, "incluida": f"D{k+1}" in seleccion, "peso": 25})
        for k in range(2):
            columnas.append({"nombre": f"Área Trabajo {k+1}", "tipo": "area",
                             "pos": k, "incluida": f"A{k+1}" in seleccion, "peso": 25})
        for k in range(2):
            columnas.append({"nombre": ["Parcial", "Recuperatorio"][k], "tipo": "otra",
                             "pos": k, "incluida": f"O{k+1}" in seleccion, "peso": 25})
        return ColumnConfig(modo=modo, columnas=[
            dict(c, incluida=bool(c["incluida"])) for c in columnas
        ])

    def _f(self, config, col_otras_start=8):
        # Periodo 3 => ev en C(3),D(4); área arranca en E(5) (col_at1=5);
        # "otras" arrancan en H(8); r=11.
        return generador._formula_definitiva(3, 5, 11, 5, config, col_otras_start)

    def test_simple_con_ev_area_y_otra(self):
        # D2 -> D(4); A1 -> E(5); O2 -> H(8)+1 = I(9).
        cfg = self._cfg({"D2", "A1", "O2"})
        self.assertEqual(self._f(cfg), '=IFERROR(AVERAGE(D11,E11,I11),"")')

    def test_pesos_mixto_con_otras(self):
        cfg = self._cfg({"D1", "A2", "O1", "O2"}, modo="pesos")
        # C (ev1), F (área2), H (otra1), I (otra2).
        self.assertEqual(
            self._f(cfg),
            '=IFERROR((C11*25+F11*25+H11*25+I11*25)/100,"")',
        )

    def test_otra_sin_col_otras_start_lanza_valueerror(self):
        cfg = self._cfg({"O1"})
        with self.assertRaises(ValueError):
            generador._formula_definitiva(3, 5, 11, 5, cfg, None)

    def test_legacy_ignora_col_otras_start(self):
        # Sin config el parámetro extra no cambia la fórmula histórica.
        self.assertEqual(
            generador._formula_definitiva(3, 5, 11, 5, None, 8),
            '=IFERROR(SUM(E11:I11)/5,"")',
        )


class TestExcelEndToEndConOtras(unittest.TestCase):
    """El Excel final incluye la zona de columnas "otras" y sus fórmulas."""

    def test_periodo_3_con_otras(self):
        # Periodo 3 => ev en C,D; 2 áreas E,F; otras G,H; definitiva en I(9).
        # El alumno 1 trae una sola nota "otra" pero el encabezado declara 2:
        # la zona queda de 2 columnas (la segunda celda en blanco).
        planilla = _planilla_con_otras(
            [[40, 50], [41, 51]],
            [[42], [38, 40]],
            n_area_trabajo=2,
            otras_columnas=["Parcial", "Trabajo Práctico"],
        )
        cfg = ColumnConfig.crear_desde_planilla(
            2, n_ev=2, nombres_otras=["Parcial", "Trabajo Práctico"]
        )
        ws = _hoja_cargada(planilla, cfg)
        hr = _fila_header(ws)
        fila1 = hr + 1
        # Títulos: G = Parcial, H = Trabajo Práctico, I = Definitiva.
        self.assertEqual(ws.cell(row=hr, column=7).value, "Parcial")
        self.assertEqual(ws.cell(row=hr, column=8).value, "Trabajo Práctico")
        self.assertEqual(ws.cell(row=hr, column=9).value, "Definitiva Periodo 3")
        # Valores: G=42, H en blanco (None) por el padding declarado.
        self.assertEqual(ws.cell(row=fila1, column=7).value, 42)
        self.assertIsNone(ws.cell(row=fila1, column=8).value)
        # Simple con todas (ev+áreas+otras): AVERAGE de C,D,E,F,G,H.
        self.assertEqual(
            ws.cell(row=fila1, column=9).value,
            f'=IFERROR(AVERAGE(C{fila1},D{fila1},E{fila1},F{fila1},G{fila1},H{fila1}),"")',
        )

    def test_periodo_4_con_otras_y_anual(self):
        # Periodo 4 => 3 evs C,D,E; 2 áreas F,G; 2 otras H,I; def J(10); anual K(11).
        planilla = _planilla_con_otras(
            [[40, 50], [41, 51]],
            [[42, 48], [38, 40]],
            n_area_trabajo=2,
            otras_columnas=["Parcial", "Definitiva 1"],
            periodo=4,
        )
        cfg = ColumnConfig(columnas=[
            {"nombre": "Def. Periodo 3", "tipo": "ev", "pos": 2, "incluida": True, "peso": 25},
            {"nombre": "Área Trabajo 2", "tipo": "area", "pos": 1, "incluida": True, "peso": 25},
            {"nombre": "Parcial", "tipo": "otra", "pos": 0, "incluida": True, "peso": 25},
            {"nombre": "Definitiva 1", "tipo": "otra", "pos": 1, "incluida": True, "peso": 25},
        ])
        ws = _hoja_cargada(planilla, cfg)
        hr = _fila_header(ws)
        fila1 = hr + 1
        # Títulos: H = Parcial, I = Definitiva 1, J = Definitiva Periodo 4, K = Anual.
        self.assertEqual(ws.cell(row=hr, column=8).value, "Parcial")
        self.assertEqual(ws.cell(row=hr, column=9).value, "Definitiva 1")
        self.assertEqual(ws.cell(row=hr, column=10).value, "Definitiva Periodo 4")
        self.assertEqual(ws.cell(row=hr, column=11).value, "Definitiva Anual")
        # Simple con la selección mixta: E (ev3), G (área2), H, I.
        self.assertEqual(
            ws.cell(row=fila1, column=10).value,
            f'=IFERROR(AVERAGE(E{fila1},G{fila1},H{fila1},I{fila1}),"")',
        )
        # Anual: promedio de las 3 evs + la definitiva del periodo.
        self.assertEqual(
            ws.cell(row=fila1, column=11).value,
            f'=IFERROR(AVERAGE(C{fila1},D{fila1},E{fila1},J{fila1}),"")',
        )
        # Valor de una celda "otra" escrita.
        self.assertEqual(ws.cell(row=fila1, column=8).value, 42)

    def test_sin_titulos_usa_nombres_genericos_y_legacy(self):
        # Sin la clave "otras_columnas" (planillas pre-cambio) pero con valores:
        # títulos genéricos "Otras notas 1..N"; sin config -> fórmula legacy de
        # áreas (las otras se escriben pero no entran al cálculo).
        planilla = _planilla_con_otras(
            [[40, 50], [41, 51]],
            [[42], [38]],
            n_area_trabajo=2,
            otras_columnas=None,
        )
        ws = _hoja_cargada(planilla, None)
        hr = _fila_header(ws)
        fila1 = hr + 1
        # ev C,D; áreas E,F; otras G; def H(8). Legacy: SUM(E:F)/2.
        self.assertEqual(ws.cell(row=hr, column=7).value, "Otras notas 1")
        self.assertEqual(ws.cell(row=fila1, column=7).value, 42)
        self.assertEqual(ws.cell(row=fila1, column=8).value,
                         f'=IFERROR(SUM(E{fila1}:F{fila1})/2,"")')


if __name__ == "__main__":
    unittest.main()